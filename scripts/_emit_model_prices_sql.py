"""Helper that converts model_prices.xlsx into plain SQL on stdout.

Used when the loader cannot reach PostgreSQL directly (e.g. a host-installed
Postgres is shadowing the docker container's published port). Pipe the output
to ``docker exec -i fluiq-postgres psql -U fluiq -d fluiq``.
"""
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

COLUMNS = [
    "provider",
    "model",
    "modality",
    "input_token_cost_per_million",
    "cached_input_token_cost_per_million",
    "output_token_cost_per_million",
    "long_context_consider_token_greater_than",
    "long_context_input_per_million",
    "long_context_cached_input_per_million",
    "long_context_output_per_million",
    "cache_valid_5_minutes_per_million",
    "cache_valid_60_minutes_per_million",
    "training_cost_per_hour",
    "video_size",
    "video_portrait",
    "video_landscape",
    "video_price_per_second",
    "audio_price_per_second",
]


def _quote(value) -> str:
    if value is None or value == "":
        return "NULL"
    if isinstance(value, (int, float, Decimal)):
        return repr(value) if not isinstance(value, float) else f"{value!r}"
    text = str(value).replace("'", "''")
    return f"'{text}'"


def main() -> None:
    if len(sys.argv) < 2:
        print("usage: _emit_model_prices_sql.py <xlsx> [out.sql]", file=sys.stderr)
        sys.exit(2)

    out = open(sys.argv[2], "w", encoding="utf-8", newline="\n") if len(sys.argv) > 2 else sys.stdout
    wb = load_workbook(Path(sys.argv[1]), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        sys.exit(0)

    header_len = len(rows[0])
    if header_len != len(COLUMNS):
        print(
            f"-- header mismatch: expected {len(COLUMNS)} got {header_len}",
            file=sys.stderr,
        )
        sys.exit(1)

    out.write("BEGIN;\n")
    out.write("TRUNCATE TABLE model_prices RESTART IDENTITY;\n")
    cols = ", ".join(COLUMNS)
    for row in rows[1:]:
        if all(v is None or v == "" for v in row):
            continue
        if not row[0] or not row[1] or not row[2]:
            continue
        values = ", ".join(_quote(v) for v in row)
        out.write(f"INSERT INTO model_prices ({cols}) VALUES ({values});\n")
    out.write("COMMIT;\n")
    if out is not sys.stdout:
        out.close()


if __name__ == "__main__":
    main()
