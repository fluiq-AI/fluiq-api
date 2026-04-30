"""One-shot loader that imports model_prices.xlsx into PostgreSQL.

Usage:
    python scripts/load_model_prices.py "D:/ideas/FluiqAI/model_prices.xlsx"
    python scripts/load_model_prices.py path/to/file.xlsx --no-truncate

The xlsx is expected to have a single sheet whose first row matches COLUMNS
below (note: the source spreadsheet has a typo "video_potrait" which is mapped
to the canonical "video_portrait" column in the database).
"""
import argparse
import asyncio
import logging
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from db_queues.postgresql import postgres_client  # noqa: E402

logger = logging.getLogger(__name__)

COLUMNS = [
    "provider",
    "model",
    "modality",
    "input_token_cost_per_million",
    "cached_input_token_cost_per_million",
    "output_token_cost_per_million",
    "long_context_consider_token_greater_than",
    "long_context_input_per_million",
    "long_context_cached_input_per_million",
    "long_context_output_per_million",
    "cache_valid_5_minutes_per_million",
    "cache_valid_60_minutes_per_million",
    "training_cost_per_hour",
    "video_size",
    "video_portrait",
    "video_landscape",
    "video_price_per_second",
    "audio_price_per_second",
]


def _read_rows(xlsx_path: Path) -> list[tuple]:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    if len(header) != len(COLUMNS):
        raise ValueError(
            f"Expected {len(COLUMNS)} columns, found {len(header)}: {header}"
        )
    data: list[tuple] = []
    for row in rows[1:]:
        if all(v is None or v == "" for v in row):
            continue
        if not row[0] or not row[1] or not row[2]:
            continue
        data.append(tuple(row))
    return data


async def load(xlsx_path: Path, truncate: bool = True) -> int:
    await postgres_client.start()
    try:
        rows = _read_rows(xlsx_path)
        insert_sql = (
            f"INSERT INTO model_prices ({', '.join(COLUMNS)}) "
            f"VALUES ({', '.join(f'${i + 1}' for i in range(len(COLUMNS)))})"
        )
        async with postgres_client.acquire() as conn:
            if truncate:
                await conn.execute("TRUNCATE TABLE model_prices RESTART IDENTITY")
            await conn.executemany(insert_sql, rows)
        return len(rows)
    finally:
        await postgres_client.stop()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Load model_prices.xlsx into PostgreSQL",
    )
    parser.add_argument("xlsx", type=Path, help="Path to the model_prices.xlsx file")
    parser.add_argument(
        "--no-truncate",
        action="store_true",
        help="Append instead of replacing existing rows",
    )
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s"
    )
    inserted = asyncio.run(load(args.xlsx, truncate=not args.no_truncate))
    logger.info("[MODEL PRICES] inserted %d rows", inserted)


if __name__ == "__main__":
    main()
